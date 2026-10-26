"""
src/export_docx.py
Executive DOCX Export Module for PitchPerfect (ChiEAC Fellowship Week 5).
Generates publication-grade Microsoft Word (.docx) documents from Pitch Kits
with native Word heading hierarchies, styled tables, bullet formatting, and brand styling.
"""

import io
import os
import re
from typing import Optional, List, Dict, Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from pitchkit import PitchKit


# ── Color Palette Constants ────────────────────────────────────────────────

COLOR_PRIMARY = RGBColor(30, 58, 138)       # #1E3A8A - Deep Navy
COLOR_SECONDARY = RGBColor(51, 65, 85)     # #334155 - Slate
COLOR_ACCENT = RGBColor(2, 132, 199)       # #0284C7 - Cerulean Blue
COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B - Muted Gray
COLOR_DARK = RGBColor(15, 23, 42)          # #0F172A - Charcoal
COLOR_GREEN = RGBColor(22, 163, 74)        # #16A34A - Pitch-Ready Green
COLOR_AMBER = RGBColor(217, 119, 6)        # #D97706 - Strong Contender Amber


# ── XML Helper Utilities for Cell Shading & Borders ────────────────────────

def _set_cell_background(cell, hex_color: str):
    """Sets background shading color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding margins in twentieths of a point (dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _set_table_borders(table, hex_color="E2E8F0"):
    """Sets subtle light-gray borders on a docx table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)


# ── Text Formatting Helpers ────────────────────────────────────────────────

def _add_runs_with_inline_formatting(paragraph, text: str, default_color=None):
    """Parses **bold** and *italic* markdown-style formatting into DOCX runs."""
    # Normalize multiple asterisks and clean special dashes
    cleaned = text.replace("****", "**").replace("–", "-").replace("—", " - ")

    bold_tokens = cleaned.split("**")
    for i, token in enumerate(bold_tokens):
        if not token:
            continue
        is_bold = (i % 2 == 1)

        if "*" in token:
            italic_tokens = token.split("*")
            for j, itok in enumerate(italic_tokens):
                if not itok:
                    continue
                run = paragraph.add_run(itok)
                if is_bold:
                    run.bold = True
                if j % 2 == 1:
                    run.italic = True
                if default_color:
                    run.font.color.rgb = default_color
        else:
            run = paragraph.add_run(token)
            if is_bold:
                run.bold = True
            if default_color:
                run.font.color.rgb = default_color


def _add_styled_heading(doc, text: str, level: int):
    """Adds a native Word Heading paragraph with consistent brand styling."""
    heading = doc.add_heading(text, level=level)
    run = heading.runs[0] if heading.runs else heading.add_run(text)
    run.font.name = "Arial"
    if level == 1:
        run.font.size = Pt(13)
        run.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        heading.paragraph_format.space_before = Pt(14)
        heading.paragraph_format.space_after = Pt(4)
    elif level == 2:
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(3)
    return heading


def _insert_evidence_table(doc, top_evidence: List[Dict[str, Any]]):
    """Constructs a 4-column executive evidence table inside the Word document."""
    if not top_evidence:
        return

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(2.5), Inches(1.3), Inches(1.4), Inches(1.8)]
    hdr_row = table.rows[0].cells
    hdr_titles = ["Evidence Metric", "Verified Value", "Sample / Period", "Verification Basis"]

    for idx, (cell, title, width) in enumerate(zip(hdr_row, hdr_titles, col_widths)):
        cell.width = width
        _set_cell_background(cell, "1E3A8A")
        _set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, ev in enumerate(top_evidence):
        row = table.add_row().cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"

        row[0].width = col_widths[0]
        row[1].width = col_widths[1]
        row[2].width = col_widths[2]
        row[3].width = col_widths[3]

        for c in row:
            _set_cell_background(c, bg_color)
            _set_cell_margins(c, top=80, bottom=80, left=100, right=100)

        # Label
        p0 = row[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(ev.get("label", ""))
        r0.font.size = Pt(9)
        r0.bold = True
        r0.font.color.rgb = COLOR_DARK

        # Formatted Value
        p1 = row[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(str(ev.get("formatted_value", "")).replace("–", "-"))
        r1.font.size = Pt(9.5)
        r1.bold = True
        r1.font.color.rgb = COLOR_PRIMARY

        # Sample & Range
        p2 = row[2].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        sample_str = f"n={ev.get('sample_size', 'N/A')}\n{ev.get('date_range', '')}"
        r2 = p2.add_run(sample_str)
        r2.font.size = Pt(8)
        r2.font.color.rgb = COLOR_MUTED

        # Method Basis
        p3 = row[3].paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        method_txt = ev.get("method_note", "")
        if len(method_txt) > 95:
            method_txt = method_txt[:95] + "..."
        r3 = p3.add_run(method_txt.replace("–", "-"))
        r3.font.size = Pt(8)
        r3.font.color.rgb = COLOR_SECONDARY

    _set_table_borders(table, "E2E8F0")

    # Spacer after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(8)
    sp.paragraph_format.space_after = Pt(0)


# ── Document Builder ────────────────────────────────────────────────────────

def build_pitch_kit_document(pitch_kit: PitchKit) -> Document:
    """
    Constructs an executive styled Word document from PitchKit.
    Integrates clean typography, title header, scorecard callout table,
    proper section flow, integrated evidence table, and gap closure callout.
    """
    doc = Document()

    # 1. Page Margins (0.75 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        # Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        clean_opp_name = pitch_kit.opportunity_name.replace("–", "-")
        f_run = f_p.add_run(f"PitchPerfect | ChiEAC Fellowship - {clean_opp_name} Proposal")
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = COLOR_MUTED

    # 2. Document Title Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    clean_title = pitch_kit.pitch_title.replace("—", " - ").replace("–", " - ")
    t_run = title_p.add_run(clean_title)
    t_run.font.name = "Arial"
    t_run.font.size = Pt(17)
    t_run.bold = True
    t_run.font.color.rgb = COLOR_PRIMARY

    # Subtitle / Metadata Line
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    s_run = sub_p.add_run(
        f"Prepared for: {pitch_kit.buyer}   |   "
        f"Instructor: {pitch_kit.instructor_name}   |   "
        f"Date: {pitch_kit.generated_date}"
    )
    s_run.font.size = Pt(9.5)
    s_run.font.color.rgb = COLOR_MUTED
    s_run.bold = True

    # 3. Readiness Scorecard Summary Callout Box
    score_table = doc.add_table(rows=1, cols=3)
    score_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    score_table.autofit = False

    widths = [Inches(2.2), Inches(2.2), Inches(2.6)]
    hdr_cells = score_table.rows[0].cells
    for i, w in enumerate(widths):
        hdr_cells[i].width = w

    # Cell 1: Readiness Score
    _set_cell_background(hdr_cells[0], "F1F5F9")
    _set_cell_margins(hdr_cells[0], top=120, bottom=120, left=140, right=140)
    p1 = hdr_cells[0].paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run("READINESS SCORE\n")
    r1.font.size = Pt(8)
    r1.font.color.rgb = COLOR_MUTED
    r1.bold = True
    r2 = p1.add_run(f"{pitch_kit.total_score} / 100")
    r2.font.size = Pt(16)
    r2.bold = True
    r2.font.color.rgb = COLOR_GREEN if pitch_kit.total_score >= 85 else COLOR_AMBER

    # Cell 2: Match Tier
    _set_cell_background(hdr_cells[1], "F1F5F9")
    _set_cell_margins(hdr_cells[1], top=120, bottom=120, left=140, right=140)
    p2 = hdr_cells[1].paragraphs[0]
    p2.paragraph_format.space_after = Pt(2)
    r3 = p2.add_run("MATCH STATUS\n")
    r3.font.size = Pt(8)
    r3.font.color.rgb = COLOR_MUTED
    r3.bold = True
    clean_tier = pitch_kit.tier.split(" ")[0].replace("–", "-")
    r4 = p2.add_run(f"{clean_tier}")
    r4.font.size = Pt(14)
    r4.bold = True
    r4.font.color.rgb = COLOR_PRIMARY

    # Cell 3: Decision Maker Priority
    _set_cell_background(hdr_cells[2], "F1F5F9")
    _set_cell_margins(hdr_cells[2], top=120, bottom=120, left=140, right=140)
    p3 = hdr_cells[2].paragraphs[0]
    p3.paragraph_format.space_after = Pt(2)
    r5 = p3.add_run("BUYER PRIORITY\n")
    r5.font.size = Pt(8)
    r5.font.color.rgb = COLOR_MUTED
    r5.bold = True
    prio_snippet = pitch_kit.buyer_priority
    if len(prio_snippet) > 85:
        prio_snippet = prio_snippet[:85] + "..."
    r6 = p3.add_run(prio_snippet.replace("–", "-"))
    r6.font.size = Pt(8.5)
    r6.italic = True
    r6.font.color.rgb = COLOR_SECONDARY

    _set_table_borders(score_table, "CBD5E1")

    # Spacer after score card
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(10)
    sp.paragraph_format.space_after = Pt(0)

    # 4. Parse Rendered Text Narrative into Structured Word Document
    raw_lines = pitch_kit.rendered_text.split("\n")

    # Step A: Filter out all front-matter lines before the first "## 1."
    body_lines = []
    found_section_1 = False
    for line in raw_lines:
        if line.strip().startswith("## 1.") or line.strip().startswith("## 1 "):
            found_section_1 = True
        if found_section_1:
            body_lines.append(line)

    if not body_lines:
        # Fallback if numbering differs
        body_lines = [l for l in raw_lines if not l.strip().startswith("# ") and not l.strip().startswith("**Prepared")]

    # Step B: Walk through body lines and render
    skip_raw_evidence_loop = False
    table_inserted = False

    for line in body_lines:
        stripped = line.strip()

        if not stripped or stripped == "---":
            continue

        # Check for Heading 2 (e.g. ## 1. Executive Summary, ## 2. Measurable Outcomes)
        if stripped.startswith("## "):
            skip_raw_evidence_loop = False
            h_text = stripped.replace("## ", "").strip()
            _add_styled_heading(doc, h_text, level=1)

            # If this is Section 2 (Evidence & Outcomes), insert the Evidence Table!
            if ("2." in h_text or "EVIDENCE" in h_text.upper() or "OUTCOMES" in h_text.upper()) and not table_inserted:
                _insert_evidence_table(doc, pitch_kit.top_evidence)
                table_inserted = True
            continue

        # Check for Heading 3 (e.g. ### Quantified Workplace Impact)
        if stripped.startswith("### "):
            skip_raw_evidence_loop = False
            h_text = stripped.replace("### ", "").strip()
            _add_styled_heading(doc, h_text, level=2)
            continue

        # Skip the redundant raw `{% for item in top_evidence %}` bullets in Section 2
        # because the 4-column executive table already rendered them above!
        if stripped.startswith("* *Evidence Basis:*") or stripped.startswith("- *Evidence Basis:*"):
            continue
        if any(ev.get("label", "") in stripped for ev in pitch_kit.top_evidence if ev.get("label")):
            continue

        # Indented / Sub-Bullet
        if line.startswith("  * ") or line.startswith("    * ") or line.startswith("  - "):
            b_text = stripped.lstrip("*- ").strip()
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.left_indent = Inches(0.5)
            bp.paragraph_format.space_before = Pt(1)
            bp.paragraph_format.space_after = Pt(2)
            _add_runs_with_inline_formatting(bp, b_text, default_color=COLOR_SECONDARY)
            continue

        # Standard Bullet (* or -)
        if stripped.startswith("* ") or stripped.startswith("- "):
            b_text = stripped[2:].strip()
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_before = Pt(1)
            bp.paragraph_format.space_after = Pt(3)
            _add_runs_with_inline_formatting(bp, b_text, default_color=COLOR_DARK)
            continue

        # Numbered List (1. , 2. )
        if re.match(r"^\d+\.\s", stripped):
            num_match = re.match(r"^(\d+\.)\s*(.*)", stripped)
            prefix = num_match.group(1)
            b_text = num_match.group(2)
            np = doc.add_paragraph(style="List Number") if "List Number" in doc.styles else doc.add_paragraph()
            np.paragraph_format.space_before = Pt(2)
            np.paragraph_format.space_after = Pt(3)
            if "List Number" not in doc.styles:
                r_num = np.add_run(f"{prefix} ")
                r_num.bold = True
            _add_runs_with_inline_formatting(np, b_text, default_color=COLOR_DARK)
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        _add_runs_with_inline_formatting(p, stripped, default_color=COLOR_DARK)

    # 5. Gap Closure Callout Box if score < 85
    if pitch_kit.gap_closure_note and pitch_kit.total_score < 85.0:
        gap_tbl = doc.add_table(rows=1, cols=1)
        gap_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        gap_cell = gap_tbl.rows[0].cells[0]
        gap_cell.width = Inches(7.0)
        _set_cell_background(gap_cell, "FFFBEB")  # Soft Amber background
        _set_cell_margins(gap_cell, top=140, bottom=140, left=160, right=160)

        gp = gap_cell.paragraphs[0]
        gp.paragraph_format.space_after = Pt(4)
        gr1 = gp.add_run("EVIDENCE BRIDGE & QUALITY ASSURANCE PROTOCOL\n")
        gr1.font.size = Pt(9.5)
        gr1.bold = True
        gr1.font.color.rgb = COLOR_AMBER

        _add_runs_with_inline_formatting(gp, pitch_kit.gap_closure_note, default_color=COLOR_DARK)
        _set_table_borders(gap_tbl, "FDE68A")

    return doc


# ── Public Exporters ────────────────────────────────────────────────────────

def export_pitch_kit_to_docx(pitch_kit: PitchKit, output_path: Optional[str] = None) -> bytes:
    """
    Exports a PitchKit instance to Microsoft Word (.docx) format.
    Returns bytes content and optionally saves to output_path if provided.
    """
    doc = build_pitch_kit_document(pitch_kit)
    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    if output_path:
        out_dir = os.path.dirname(os.path.abspath(output_path))
        os.makedirs(out_dir, exist_ok=True)
        with open(output_path, "wb") as f:
            f.write(docx_bytes)

    return docx_bytes


def save_pitch_kit_docx(pitch_kit: PitchKit, file_path: str) -> str:
    """Convenience helper to save pitch kit directly to a docx file."""
    export_pitch_kit_to_docx(pitch_kit, output_path=file_path)
    return file_path
