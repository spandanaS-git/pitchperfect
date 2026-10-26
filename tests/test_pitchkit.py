"""
tests/test_pitchkit.py
Comprehensive unit tests for PitchPerfect Week 5:
Pitch Kit Generator, Jinja2 Templates, and DOCX Export.
"""

import io
import os
import sys
import pytest
import pandas as pd
from docx import Document

# Add src to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src")))

from evidence import compute_all_evidence_cards
from scorecard import calculate_opportunity_scorecard, load_taxonomy
from pitchkit import (
    PitchKit,
    load_template,
    get_top_evidence_for_opportunity,
    generate_gap_closure_note,
    generate_pitch_kit,
    generate_all_pitch_kits,
    export_pitch_kit_to_text,
    export_pitch_kit_to_markdown
)
from export_docx import export_pitch_kit_to_docx, build_pitch_kit_document


OPPORTUNITY_IDS = [
    "studio_residency",
    "corporate_wellness",
    "community_organization",
    "school",
    "festival",
    "hotel",
    "private_event"
]


@pytest.fixture
def demo_cards():
    """Loads demo CSV files and computes verified Evidence Cards."""
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    demo_dir = os.path.join(base_dir, "data", "demo")

    c_df = pd.read_csv(os.path.join(demo_dir, "class_history.csv"))
    s_df = pd.read_csv(os.path.join(demo_dir, "survey_responses.csv"))
    e_df = pd.read_csv(os.path.join(demo_dir, "event_outcomes.csv"))

    return compute_all_evidence_cards(c_df, s_df, e_df)


@pytest.fixture
def demo_scorecard(demo_cards):
    """Calculates full 7-opportunity scorecard for demo data."""
    return calculate_opportunity_scorecard(demo_cards)


# ── Test Suite ─────────────────────────────────────────────────────────────

def test_load_all_7_templates():
    """Verifies that all 7 Jinja2 opportunity templates compile cleanly."""
    for opp_id in OPPORTUNITY_IDS:
        template = load_template(opp_id)
        assert template is not None, f"Failed to load template for {opp_id}"


def test_missing_template_raises_error():
    """Verifies that attempting to load a non-existent template raises FileNotFoundError."""
    with pytest.raises(FileNotFoundError):
        load_template("invalid_opportunity_xyz")


def test_generate_pitch_kit_studio_residency(demo_cards, demo_scorecard):
    """Tests generating a pitch kit for Studio Residency with custom profile."""
    profile = {
        "name": "Mariana Rodriguez",
        "disciplines": "Vinyasa Yoga & Breathwork",
        "years_teaching": 5,
        "email": "mariana@example.com",
        "phone": "(312) 555-0144"
    }
    kit = generate_pitch_kit("studio_residency", demo_cards, demo_scorecard, instructor_profile=profile)

    assert isinstance(kit, PitchKit)
    assert kit.opportunity_id == "studio_residency"
    assert kit.opportunity_name == "Studio Residency"
    assert kit.instructor_name == "Mariana Rodriguez"
    assert "Studio Owner" in kit.buyer
    assert kit.total_score > 0
    assert len(kit.top_evidence) > 0
    assert len(kit.rendered_text) > 300
    assert "Mariana Rodriguez" in kit.rendered_text
    assert "Studio Residency" in kit.rendered_text


def test_generate_pitch_kit_all_7_opportunities(demo_cards, demo_scorecard):
    """Verifies generating pitch kits across all 7 opportunity types."""
    profile = {"name": "Test Instructor", "years_teaching": 4}
    for opp_id in OPPORTUNITY_IDS:
        kit = generate_pitch_kit(opp_id, demo_cards, demo_scorecard, instructor_profile=profile)
        assert isinstance(kit, PitchKit)
        assert kit.opportunity_id == opp_id
        assert kit.instructor_name == "Test Instructor"
        assert kit.buyer != ""
        assert kit.buyer_priority != ""
        assert len(kit.rendered_text) > 200
        assert "Test Instructor" in kit.rendered_text


def test_pitch_kit_tone_and_narrative_distinctness(demo_cards, demo_scorecard):
    """
    Verifies that each opportunity pitch has distinct vocabulary, tone,
    and narrative focus aligned with its buyer's decision model.
    """
    profile = {"name": "Elena Smith", "years_teaching": 6}

    # 1. Corporate Wellness: ROI, stress reduction, workday midday fit
    corp_kit = generate_pitch_kit("corporate_wellness", demo_cards, demo_scorecard, instructor_profile=profile)
    c_text = corp_kit.rendered_text.lower()
    assert any(w in c_text for w in ["stress", "workplace", "lunch", "midday", "employee"])

    # 2. School: Safety, background check, youth, classroom
    school_kit = generate_pitch_kit("school", demo_cards, demo_scorecard, instructor_profile=profile)
    s_text = school_kit.rendered_text.lower()
    assert any(w in s_text for w in ["youth", "background check", "student", "classroom", "safety"])

    # 3. Hotel: 5-star, sunrise/morning, luxury, guest satisfaction
    hotel_kit = generate_pitch_kit("hotel", demo_cards, demo_scorecard, instructor_profile=profile)
    h_text = hotel_kit.rendered_text.lower()
    assert any(w in h_text for w in ["guest", "sunrise", "morning", "hospitality", "luxury", "5-star"])

    # 4. Festival: Crowd, stage, outdoor, energy
    fest_kit = generate_pitch_kit("festival", demo_cards, demo_scorecard, instructor_profile=profile)
    f_text = fest_kit.rendered_text.lower()
    assert any(w in f_text for w in ["crowd", "stage", "outdoor", "festival", "energy"])

    # 5. Community Organization: Belonging, welcomed, sliding scale, neighborhood
    comm_kit = generate_pitch_kit("community_organization", demo_cards, demo_scorecard, instructor_profile=profile)
    cm_text = comm_kit.rendered_text.lower()
    assert any(w in cm_text for w in ["belonging", "welcomed", "sliding-scale", "community", "equity"])

    # 6. Studio Residency: Fill rate, retention, weekly slot
    studio_kit = generate_pitch_kit("studio_residency", demo_cards, demo_scorecard, instructor_profile=profile)
    st_text = studio_kit.rendered_text.lower()
    assert any(w in st_text for w in ["fill rate", "retention", "residency", "weekly"])

    # 7. Private Event: Bespoke, retreat, rebooking, rate card
    priv_kit = generate_pitch_kit("private_event", demo_cards, demo_scorecard, instructor_profile=profile)
    p_text = priv_kit.rendered_text.lower()
    assert any(w in p_text for w in ["bespoke", "retreat", "rebooking", "rate card", "private"])


def test_instructor_name_default_and_custom(demo_cards, demo_scorecard):
    """Verifies instructor name falls back to 'Your Name' or uses provided value."""
    # Default fallback
    default_kit = generate_pitch_kit("studio_residency", demo_cards, demo_scorecard, instructor_profile={})
    assert default_kit.instructor_name == "Your Name"
    assert "Your Name" in default_kit.rendered_text

    # Custom name
    custom_kit = generate_pitch_kit(
        "studio_residency", demo_cards, demo_scorecard,
        instructor_profile={"name": "Alex Chen"}
    )
    assert custom_kit.instructor_name == "Alex Chen"
    assert "Alex Chen" in custom_kit.rendered_text


def test_gap_closure_note_logic(demo_cards):
    """Verifies gap closure note appears when total score is below 85, and is absent when >= 85."""
    # Create an artificially low-scoring scorecard by turning off all policy credentials
    all_false_policies = {
        "liability_insurance_on_file": False,
        "background_check_on_file": False,
        "sliding_scale_policy": False,
        "language_matches": False,
        "professional_certifications": False,
        "clear_rate_card_on_file": False,
        "park_community_venue_history": False,
        "youth_segment_experience": False,
        "outdoor_event_track_record": False,
        "private_event_experience": False,
    }
    low_scorecard = calculate_opportunity_scorecard(
        [],
        policy_overrides=all_false_policies
    )

    low_kit = generate_pitch_kit("corporate_wellness", [], low_scorecard)
    assert low_kit.total_score < 85.0
    assert low_kit.gap_closure_note is not None
    assert "Evidence Bridge" in low_kit.rendered_text

    # With high score (>= 85)
    high_scorecard = calculate_opportunity_scorecard(demo_cards)
    # Check top opportunity
    top_id = high_scorecard.top_opportunity_id
    top_score = high_scorecard.opportunity_scores[top_id].total_score
    if top_score >= 85.0:
        high_kit = generate_pitch_kit(top_id, demo_cards, high_scorecard)
        assert high_kit.gap_closure_note is None


def test_generate_all_pitch_kits(demo_cards, demo_scorecard):
    """Verifies batch generating pitch kits for all 7 opportunities."""
    all_kits = generate_all_pitch_kits(demo_cards, demo_scorecard)
    assert isinstance(all_kits, dict)
    assert len(all_kits) == 7
    for opp_id in OPPORTUNITY_IDS:
        assert opp_id in all_kits
        assert isinstance(all_kits[opp_id], PitchKit)


def test_export_pitch_kit_to_docx(demo_cards, demo_scorecard):
    """Verifies generating and validating a publication-grade DOCX document."""
    kit = generate_pitch_kit("corporate_wellness", demo_cards, demo_scorecard)
    docx_bytes = export_pitch_kit_to_docx(kit)

    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 2000

    # Read back with python-docx to verify valid OpenXML structure
    doc = Document(io.BytesIO(docx_bytes))
    assert len(doc.paragraphs) > 5
    assert len(doc.tables) >= 1

    # Check that document text contains core kit elements
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "CORPORATE WELLNESS" in full_text.upper()
    assert kit.buyer.upper() in full_text.upper()


def test_top_evidence_ranking(demo_cards, demo_scorecard):
    """Verifies get_top_evidence_for_opportunity selects relevant weighted cards."""
    studio_ev = get_top_evidence_for_opportunity("studio_residency", demo_cards, limit=4)
    assert len(studio_ev) <= 4
    assert len(studio_ev) > 0
    # studio residency prioritizes fill rate or retention
    ev_ids = [e["id"] for e in studio_ev]
    assert any("fill_rate" in eid or "attendance" in eid or "conversion" in eid for eid in ev_ids)


def test_empty_evidence_cards_graceful_fallback():
    """Verifies that pitch kits can be generated even if evidence cards are empty."""
    kit = generate_pitch_kit("hotel", [])
    assert isinstance(kit, PitchKit)
    assert kit.opportunity_id == "hotel"
    assert len(kit.rendered_text) > 100
    assert kit.instructor_name == "Your Name"
